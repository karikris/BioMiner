"""Contract tests for the aggregate-only Ground Zero EDA report."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import subprocess

import polars as pl
import pytest
from PIL import Image
from pptx import Presentation

from biominer.reports.ground_zero_eda import build_ground_zero_eda_run


def _fixture_source(path: Path) -> Path:
    """Write a deliberately string-typed, GBIF-shaped occurrence fixture."""
    frame = pl.DataFrame(
        {
            "occurrenceID": ["occ-1", "occ-2", "occ-3", "occ-4"],
            "year": ["2024", "2024", "", "=2022"],
            "month": ["1", None, "12", "2"],
            "decimalLatitude": ["12.9716", "", "-33.8688", "@bad"],
            "decimalLongitude": ["77.5946", "77.0", "151.2093", "151.0"],
            "kingdom": ["Animalia"] * 4,
            "phylum": ["Arthropoda"] * 4,
            "class": ["Insecta"] * 4,
            "order": ["Lepidoptera"] * 4,
            "family": ["Papilionidae", "Papilionidae", "Nymphalidae", ""],
            "genus": ["Papilio", "Papilio", "Danaus", "=Formula"],
            "species": ["Papilio demoleus", "Papilio demoleus", "Danaus plexippus", ""],
            "scientificName": ["Papilio demoleus", "Papilio demoleus", "Danaus plexippus", "=formula"],
            "taxonRank": ["SPECIES", "SPECIES", "SPECIES", ""],
            "datasetName": ["Citizen science", "Citizen science", "+Museum", "=Formula data"],
            "publisher": ["Publisher A", "Publisher A", "Publisher B", "=Publisher C"],
            "institutionCode": ["INST", "INST", "MUS", ""],
            "recordedBy": ["Alice", "Alice", "Bob", "@collector"],
            "basisOfRecord": ["HUMAN_OBSERVATION", "HUMAN_OBSERVATION", "-SPECIMEN", ""],
            "issue": ["", "COORDINATE_INVALID; TAXON_MATCH_NONE", "TAXON_MATCH_NONE", ""],
            "coordinateUncertaintyInMeters": ["10", "", "100", ""],
            "identifiedBy": ["Alice", "Alice", "Curator", ""],
        },
        strict=False,
    ).with_columns(pl.all().cast(pl.String))
    frame.write_parquet(path)
    return path


def test_ground_zero_eda_creates_safe_aggregate_report(tmp_path: Path) -> None:
    source = _fixture_source(tmp_path / "occurrences.parquet")
    source_manifest = tmp_path / "source_manifest.json"
    source_manifest.write_text(json.dumps({"source": "fixture", "version": 1}), encoding="utf-8")
    output = tmp_path / "eda"
    source_sha256_before = hashlib.sha256(source.read_bytes()).hexdigest()

    result = build_ground_zero_eda_run(source, output, source_manifest, top_n=20)

    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_sha256_before

    assert result["output_dir"] == str(output)
    work_path = output / "eda_work.parquet"
    assert work_path.exists()
    work = pl.read_parquet(work_path)
    assert {"section", "metric", "dimension", "value", "record_count"} <= set(work.columns)
    assert "occurrenceID" not in work.columns
    assert "(NULL)" in work.filter(pl.col("metric") == "occurrences_by_month")["dimension"].to_list()
    issue_counts = dict(
        work.filter(pl.col("metric") == "occurrences_by_issue").select("dimension", "value").iter_rows()
    )
    assert issue_counts == {"(NULL)": 2, "TAXON_MATCH_NONE": 2, "COORDINATE_INVALID": 1}
    assert "COORDINATE_INVALID; TAXON_MATCH_NONE" not in issue_counts

    csv_paths = sorted(output.glob("*.csv"))
    assert csv_paths
    for csv_path in csv_paths:
        raw = csv_path.read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf")
        assert raw.decode("utf-8-sig")
    combined_csv = "\n".join(path.read_text(encoding="utf-8-sig") for path in csv_paths)
    assert "'=Formula" in combined_csv
    assert "'+Museum" in combined_csv
    assert "'-SPECIMEN" in combined_csv
    assert "'@collector" in combined_csv

    charts = sorted(output.glob("*.png"))
    metric_pairs = set(work.select("section", "metric").unique().iter_rows())
    expected_chart_names = {f"{section}__{metric}.png" for section, metric in metric_pairs}
    assert {chart.name for chart in charts} == expected_chart_names
    for chart in charts:
        assert chart.stat().st_size > 0
        with Image.open(chart) as image:
            image.verify()

    deck_path = output / "ground_zero_eda.pptx"
    assert deck_path.exists() and deck_path.stat().st_size > 0
    assert len(Presentation(deck_path).slides) == 1 + len(metric_pairs)

    manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["source"]["sha256"] == source_sha256_before
    assert manifest["source"]["schema"]["occurrenceID"] == "string"
    assert "publishingOrgKey" not in manifest["source"]["schema"]
    assert manifest["source_manifest"] == json.loads(source_manifest.read_text(encoding="utf-8"))
    inventory = manifest["artifact_inventory"]
    assert inventory and all(entry["sha256"] and entry["path"] != "manifest.json" for entry in inventory)
    assert {entry["path"] for entry in inventory} >= {"eda_work.parquet", "ground_zero_eda.pptx"}
    for entry in inventory:
        assert entry["sha256"] == hashlib.sha256((output / entry["path"]).read_bytes()).hexdigest()
    assert set(manifest["artifacts"]["charts"]) == expected_chart_names
    assert not (output / "duckdb_spill").exists()
    assert manifest["query_strategy"] == {
        "engine": "duckdb",
        "input": "read-only parquet_scan",
        "completeness": "single wide aggregate scan",
        "spill": {"temporary_directory": "duckdb_spill", "max_size": "2GiB", "cleaned": True},
    }


def test_ground_zero_eda_cli_prints_manifest_json(tmp_path: Path) -> None:
    source = _fixture_source(tmp_path / "occurrences.parquet")
    source_manifest = tmp_path / "source_manifest.json"
    source_manifest.write_text(json.dumps({"source": "fixture", "version": 1}), encoding="utf-8")
    output = tmp_path / "eda"
    repository_root = Path(__file__).resolve().parents[1]

    completed = subprocess.run(
        [
            repository_root / ".venv" / "bin" / "python",
            "scripts/run_ground_zero_eda.py",
            "--source",
            str(source),
            "--source-manifest",
            str(source_manifest),
            "--output",
            str(output),
            "--top-n",
            "2",
        ],
        cwd=repository_root,
        capture_output=True,
        check=False,
        text=True,
    )

    assert completed.returncode == 0, completed.stderr
    reported = json.loads(completed.stdout)
    assert reported["manifest_path"] == str(output / "manifest.json")
    assert reported["manifest"] == json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    assert reported["manifest"]["source_manifest"] == json.loads(source_manifest.read_text(encoding="utf-8"))
    assert reported["manifest"]["top_n"] == 2


def test_ground_zero_eda_cli_reports_existing_output_as_json_error(tmp_path: Path) -> None:
    source = _fixture_source(tmp_path / "occurrences.parquet")
    output = tmp_path / "already-exists"
    output.mkdir()
    repository_root = Path(__file__).resolve().parents[1]

    completed = subprocess.run(
        [
            repository_root / ".venv" / "bin" / "python",
            "scripts/run_ground_zero_eda.py",
            "--source",
            str(source),
            "--output",
            str(output),
        ],
        cwd=repository_root,
        capture_output=True,
        check=False,
        text=True,
    )

    assert completed.returncode == 2
    assert completed.stdout == ""
    assert json.loads(completed.stderr) == {
        "error": f"refusing to overwrite existing output directory: {output}"
    }
    assert "Traceback" not in completed.stderr


def test_ground_zero_eda_cli_reports_corrupt_parquet_as_json_error(tmp_path: Path) -> None:
    source = tmp_path / "corrupt.parquet"
    source.write_text("this is not a Parquet file", encoding="utf-8")
    repository_root = Path(__file__).resolve().parents[1]

    completed = subprocess.run(
        [
            repository_root / ".venv" / "bin" / "python",
            "scripts/run_ground_zero_eda.py",
            "--source",
            str(source),
            "--output",
            str(tmp_path / "eda"),
        ],
        cwd=repository_root,
        capture_output=True,
        check=False,
        text=True,
    )

    assert completed.returncode == 2
    assert completed.stdout == ""
    assert json.loads(completed.stderr)["error"]
    assert "Traceback" not in completed.stderr


def test_ground_zero_eda_refuses_existing_output_directory(tmp_path: Path) -> None:
    source = _fixture_source(tmp_path / "occurrences.parquet")
    output = tmp_path / "already-exists"
    output.mkdir()

    with pytest.raises(FileExistsError, match="refus"):
        build_ground_zero_eda_run(source, output)


def test_ground_zero_eda_retains_null_month_beyond_top_n(tmp_path: Path) -> None:
    source = _fixture_source(tmp_path / "occurrences.parquet")
    base = pl.read_parquet(source)
    month_one = base.filter(pl.col("month") == "1")
    month_two = base.filter(pl.col("month") == "2")
    month_twelve = base.filter(pl.col("month") == "12")
    # The NULL month has one occurrence, strictly behind the two populated leaders.
    pl.concat([base, *([month_one] * 4), *([month_two] * 2), month_twelve]).write_parquet(source)

    build_ground_zero_eda_run(source, tmp_path / "eda", top_n=2)

    months = pl.read_parquet(tmp_path / "eda" / "eda_work.parquet").filter(
        pl.col("metric") == "occurrences_by_month"
    )
    assert months["dimension"].to_list() == ["1", "2", "(NULL)"]
    assert months.filter(pl.col("dimension") == "(NULL)")["value"].item() == 1


def test_ground_zero_eda_requires_only_fields_used_by_summary_sql(tmp_path: Path) -> None:
    source = _fixture_source(tmp_path / "occurrences.parquet")
    unused = [
        "occurrenceID", "kingdom", "phylum", "class", "order", "taxonRank",
        "coordinateUncertaintyInMeters",
    ]
    pl.read_parquet(source).drop(unused).write_parquet(source)

    build_ground_zero_eda_run(source, tmp_path / "accepted")

    missing_required = tmp_path / "missing-issue.parquet"
    pl.read_parquet(source).drop("issue").write_parquet(missing_required)
    with pytest.raises(ValueError, match="issue"):
        build_ground_zero_eda_run(missing_required, tmp_path / "rejected")


def test_ground_zero_eda_deck_is_byte_identical_for_identical_source(tmp_path: Path) -> None:
    source = _fixture_source(tmp_path / "occurrences.parquet")

    build_ground_zero_eda_run(source, tmp_path / "first")
    build_ground_zero_eda_run(source, tmp_path / "second")

    assert (tmp_path / "first" / "ground_zero_eda.pptx").read_bytes() == (
        tmp_path / "second" / "ground_zero_eda.pptx"
    ).read_bytes()


def test_ground_zero_eda_limits_only_rendered_chart_bars(tmp_path: Path) -> None:
    source = _fixture_source(tmp_path / "occurrences.parquet")
    base_row = pl.read_parquet(source).head(1)
    rows = [
        base_row.with_columns(
            pl.lit(f"Dataset with a deliberately long label number {index:02d}").alias("datasetName"),
            pl.lit(f"Publisher with a deliberately long label number {index:02d}").alias("publisher"),
        )
        for index in range(20)
    ]
    pl.concat(rows).write_parquet(source)

    build_ground_zero_eda_run(source, tmp_path / "eda", top_n=20)

    work = pl.read_parquet(tmp_path / "eda" / "eda_work.parquet")
    assert work.filter(pl.col("metric") == "occurrences_by_dataset").height == 20
    assert work.filter(pl.col("metric") == "occurrences_by_publisher").height == 20
    manifest = json.loads((tmp_path / "eda" / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["chart_policy"] == {"max_bars": 12, "label_wrap_width": 26, "max_label_lines": 2}
    metadata = {entry["metric"]: entry for entry in manifest["chart_metadata"]}
    assert metadata["occurrences_by_dataset"]["available_ranked_values"] == 20
    assert metadata["occurrences_by_dataset"]["displayed_bars"] == 12
    assert metadata["occurrences_by_publisher"]["available_ranked_values"] == 20
    assert metadata["occurrences_by_publisher"]["displayed_bars"] == 12
